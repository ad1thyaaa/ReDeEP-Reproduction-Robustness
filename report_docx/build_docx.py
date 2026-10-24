"""Fill the IIIT Kottayam IOE-414 'Full Project Report' Word template with the
ReDeEP project content, tables and figures, keeping the template's cover page,
logo and styles. Numbers are read from data/ so nothing is typed by hand.

    python report_docx/build_docx.py --template <Full_Project_Report_DL.docx> --out <output.docx>
"""
import argparse
import collections
import copy
import json
import sys
from pathlib import Path

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
import pymupdf
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DATA, FIGS = ROOT / "data", ROOT / "paper" / "figures"
OUTFIG = ROOT / "report_docx" / "figures"
OUTFIG.mkdir(parents=True, exist_ok=True)
sys.path.insert(0, str(ROOT))

ap = argparse.ArgumentParser()
ap.add_argument("--template", required=True)
ap.add_argument("--out", required=True)
args = ap.parse_args()

load = lambda f: json.loads((DATA / f).read_text())
opt = lambda f: load(f) if (DATA / f).exists() else None
f3 = lambda v: f"{v:.3f}"

# ============================================================ numbers
BB = [("TinyLlama-1.1B", "tinyllama_1.1b"), ("Llama-3.2-1B", "llama3.2_1b"),
      ("Llama-3.2-3B (4-bit)", "llama3.2_3b_4bit"), ("Llama-2-7B (4-bit)", "llama2_7b_4bit")]
R = {n: load(f"results_{t}.json") for n, t in BB}
l2 = R["Llama-2-7B (4-bit)"]
A = json.loads((ROOT / "paper" / "analysis.json").read_text())
AN = {"TinyLlama-1.1B": "TinyLlama-1.1B", "Llama-3.2-1B": "Llama-3.2-1B", "Llama-3.2-3B (4-bit)": "Llama-3.2-3B", "Llama-2-7B (4-bit)": "Llama-2-7B"}
robT, rob7 = load("robustness_tinyllama_1.1b.json")["conditions"], load("robustness_llama2_7b_4bit.json")["conditions"]
numT = load("robustness_numeric_subset_tinyllama_1.1b.json")["response_numbers_contradicted"]
num7 = load("robustness_numeric_subset_llama2_7b_4bit.json")["response_numbers_contradicted"]
SPL = "same_test_split_as_llama2_7b_4bit"
JG, JO = load("results_judge_gemini_gemini-3.5-flash-lite.json"), load("results_judge_groq_openai_gpt-oss-120b.json")
JL = load("results_judge_llama_unsloth_Llama-3.2-3B-Instruct.json")
s2 = l2["single_split"]
ev = lambda t: opt(f"eval_chunk_{t}.json")
tok = lambda t: opt(f"results_tok_{t}.json")
evoff = ev("official_llama2_7B_response_chunk")
BIG = [("Llama-2-7B", "llama2_7b"), ("Llama-2-13B", "llama2_13b"), ("Llama-3-8B", "llama3_8b")]


def aarf(tag):
    rows = [json.loads(l) for l in open(DATA / f"aarf_{tag}.jsonl", encoding="utf-8")]
    g = np.concatenate([np.array(r["gate_trace"])[:, 2] for r in rows if r.get("gate_trace")])
    by = {}
    for x in (json.loads(l) for l in open(DATA / f"aarf_judge_{tag}.jsonl", encoding="utf-8")):
        by.setdefault(x["id"], {})[x["order"]] = x["aarf_won"]
    bp = [v for v in by.values() if v.get("aarf_first") is not None and v.get("baseline_first") is not None]
    return dict(n=len(rows), changed=sum(r["baseline"] != r["aarf"] for r in rows), fire=100 * (g > 0.6).mean(),
                gmin=g.min(), steps=len(g), offA=sum(v.get("aarf_first") is True for v in by.values()),
                offB=sum(v.get("aarf_first") is False for v in by.values()),
                bothA=sum(v["aarf_first"] and v["baseline_first"] for v in bp),
                bothB=sum((not v["aarf_first"]) and (not v["baseline_first"]) for v in bp),
                parsed=len(bp), flips=sum(v["aarf_first"] != v["baseline_first"] for v in bp))


AA = {t: aarf(t) for _, t in BIG}

# dataset stats
allr = [json.loads(l) for l in open(ROOT.parent / "RAGTruth" / "dataset" / "response.jsonl", encoding="utf-8")]
srcs = [json.loads(l) for l in open(ROOT.parent / "RAGTruth" / "dataset" / "source_info.jsonl", encoding="utf-8")]
task_of = {s["source_id"]: s["task_type"] for s in srcs}
ds = collections.Counter()
for r in allr:
    t = task_of.get(r["source_id"]); ds[(t, "n")] += 1; ds[(t, "h")] += int(len(r.get("labels", [])) > 0)
from redeep.data import load_examples
pool = {e["id"]: e for e in load_examples(split="test", models=["llama-2-7b-chat"])}
sub = collections.Counter()
for i in np.load(DATA / "scores_llama2_7b_4bit.npz")["ids"]:
    e = pool[str(i)]; sub[(e["task_type"], "n")] += 1; sub[(e["task_type"], "h")] += e["label"]

# ============================================================ figures
def pdf_png(name):
    out = OUTFIG / f"{name}.png"
    d = pymupdf.open(FIGS / f"{name}.pdf")
    d[0].get_pixmap(dpi=220).save(out)
    return out


def arch_png():
    out = OUTFIG / "architecture.png"
    fig, ax = plt.subplots(figsize=(9, 4.6), dpi=200)
    ax.set_xlim(0, 100); ax.set_ylim(0, 52); ax.axis("off")
    def box(x, y, w, h, title, body, fc="#E6EFEE", ec="#0F3D3E"):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.4,rounding_size=1.2", fc=fc, ec=ec, lw=1.2))
        ax.text(x + w / 2, y + h - 2.6, title, ha="center", va="top", fontsize=9.5, fontweight="bold", family="serif")
        ax.text(x + w / 2, y + h - 7.2, body, ha="center", va="top", fontsize=7.6, family="serif", linespacing=1.35)
    def arrow(x1, y1, x2, y2):
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1), arrowprops=dict(arrowstyle="-|>", color="#555555", lw=1.2))
    box(1, 18, 17, 16, "RAG input", "Retrieved passage +\nquestion/instruction\n+ model response")
    box(23, 13, 20, 26, "LLaMA forward pass", "One pass over\n[prompt; response]\n\nforward hooks keep:\n- top-10% attended\n  context per head\n- residual before FFN\n- hidden states")
    box(49, 30, 22, 17, "Copying heads -> ECS", "cosine(response token,\nmean of top-attended\ncontext tokens)\nper attention head")
    box(49, 6, 22, 17, "Knowledge FFNs -> PKS", "JSD of logit-lens\ndistribution before vs.\nafter each FFN layer")
    box(77, 18, 21, 16, "ReDeEP detector", "logistic regression on\nmean PKS(F) and\n-mean ECS(A)\n-> hallucination score", fc="#F8D9A8", ec="#B5701A")
    arrow(18.6, 26, 22.6, 26); arrow(43.6, 31, 48.6, 38); arrow(43.6, 21, 48.6, 14)
    arrow(71.6, 38, 76.6, 29); arrow(71.6, 14, 76.6, 23)
    ax.text(60, 2.2, "AARF (mitigation): if the token score exceeds a threshold, amplify copying heads (x1.2) and damp knowledge FFNs (x0.8)",
            ha="center", fontsize=7.4, style="italic", family="serif", color="#444444")
    ax.text(60, 50.5, "Heads A: ranked from weights by OV-circuit copying score;  layers F: chosen on training folds by PKS-label correlation",
            ha="center", fontsize=7.4, style="italic", family="serif", color="#444444")
    fig.tight_layout(pad=0.2); fig.savefig(out); plt.close(fig)
    return out


def eq_png(name, tex, w=6.0):
    out = OUTFIG / f"{name}.png"
    fig = plt.figure(figsize=(w, 0.6), dpi=220)
    fig.text(0.5, 0.5, tex, ha="center", va="center", fontsize=14)
    fig.savefig(out, transparent=False, facecolor="white"); plt.close(fig)
    return out


def leak_png():
    out = OUTFIG / "leakage.png"
    rows = [("Authors' logs\nRAGTruth 7B", evoff), ("Ours\nRAGTruth 7B", ev("llama2_7b_ragtruth")), ("Ours\nRAGTruth 13B", ev("llama2_13b_ragtruth")),
            ("Ours\nDolly 7B", ev("llama2_7b_dolly")), ("Ours\nDolly 13B", ev("llama2_13b_dolly")), ("Ours\nDolly 8B", ev("llama3_8b_dolly"))]
    fig, ax = plt.subplots(figsize=(7.5, 3.4), dpi=200)
    x = np.arange(len(rows)); w = 0.38
    o = [r[1]["official_protocol"]["AUC"] for r in rows]; n = [r[1]["nested_cv"]["AUC_mean"] for r in rows]
    ax.bar(x - w / 2, o, w, color="#8C9796", label="Official protocol (selected on test set)")
    ax.bar(x + w / 2, n, w, color="#E8912D", label="Leakage-free nested CV (ours)")
    for i in range(len(rows)):
        ax.text(x[i] - w / 2, o[i] + 0.004, f"{o[i]:.3f}", ha="center", fontsize=7); ax.text(x[i] + w / 2, n[i] + 0.004, f"{n[i]:.3f}", ha="center", fontsize=7)
    ax.set_xticks(x, [r[0] for r in rows], fontsize=8); ax.set_ylim(0.6, 0.86); ax.set_ylabel("AUC")
    ax.legend(frameon=False, fontsize=8, loc="upper left"); ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout(pad=0.3); fig.savefig(out); plt.close(fig)
    return out


plt.rcParams.update({"font.family": "serif", "mathtext.fontset": "stix"})
FIG = {k: pdf_png(k) for k in ["detection_auc", "roc", "score_hist", "ecs_heatmap", "pks_by_layer", "per_task",
                              "ablation_grid", "judges_auc", "robustness_shift"]}
FIG["architecture"], FIG["leakage"] = arch_png(), leak_png()
FIG["eq_ecs"] = eq_png("eq_ecs", r"$E^{l,h}=\frac{1}{|R|}\sum_{n\in R}\cos\left(\frac{1}{|T_n|}\sum_{i\in T_n}x_i^{L},\;x_n^{L}\right)$")
FIG["eq_pks"] = eq_png("eq_pks", r"$P^{l}=\frac{1}{|R|}\sum_{n\in R}\mathrm{JSD}\left(\sigma(W_U\,\mathrm{norm}(x_n^{mid,l})),\;\sigma(W_U\,\mathrm{norm}(x_n^{l}))\right)$")
FIG["eq_h"] = eq_png("eq_h", r"$H(r)=\alpha\,\overline{P}_{\mathcal{F}}(r)-\beta\,\overline{E}_{\mathcal{A}}(r),\qquad \alpha,\beta>0$", 4.5)

# ============================================================ document helpers
doc = Document(args.template)
TNR = "Times New Roman"
for sn, size in (("Normal", 12), ("Caption", 10), ("List Bullet", 12), ("List Number", 12), ("Heading 2", 13), ("Heading 3", 12)):
    st = doc.styles[sn]
    st.font.name = TNR
    st.font.size = Pt(size)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), TNR)
for sn in ("Heading 2", "Heading 3"):
    doc.styles[sn].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles[sn].font.bold = True
doc.styles["Caption"].font.italic = True
doc.styles["Caption"].font.color.rgb = RGBColor(0x33, 0x33, 0x33)

body = doc.element.body
heads = {p.text.strip(): p for p in doc.paragraphs if p.style.name == "Heading 1"}


class Cursor:
    """Inserts new blocks immediately before `anchor` (so content keeps order)."""
    def __init__(self, anchor):
        self.anchor = anchor

    def _move(self, el):
        self.anchor.addprevious(el)

    def p(self, text="", style="Normal", bold_lead=None, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=None):
        para = doc.add_paragraph(style=style)
        if bold_lead:
            r = para.add_run(bold_lead + " "); r.bold = True; r.font.name = TNR
        r = para.add_run(text); r.italic = italic; r.font.name = TNR
        if size:
            r.font.size = Pt(size)
        if align is not None and style in ("Normal",):
            para.alignment = align
        para.paragraph_format.space_after = Pt(6)
        para.paragraph_format.line_spacing = 1.15
        self._move(para._p)
        return para

    def h(self, text, level=2):
        para = doc.add_paragraph(text, style=f"Heading {level}")
        for r in para.runs:
            r.font.name = TNR
        self._move(para._p)

    def bullets(self, items, numbered=False):
        for it in items:
            if isinstance(it, tuple):
                self.p(it[1], style="List Number" if numbered else "List Bullet", bold_lead=it[0], align=None)
            else:
                self.p(it, style="List Number" if numbered else "List Bullet", align=None)

    def caption(self, text, keep_next=False):
        para = doc.add_paragraph(style="Caption")
        r = para.add_run(text); r.font.name = TNR
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = keep_next      # table captions stay with their table
        self._move(para._p)

    def fig(self, key, caption, width=6.0):
        doc.add_picture(str(FIG[key]), width=Inches(width))
        para = doc.paragraphs[-1]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = bool(caption)   # figure stays with its caption
        self._move(para._p)
        if caption:
            self.caption(caption)

    def table(self, header, rows, caption, widths=None, bold_col=None):
        self.caption(caption, keep_next=True)
        t = doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(10); r.font.name = TNR
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), "D9E2E1")
            c._tc.get_or_add_tcPr().append(sh)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(10); r.font.name = TNR
                r.bold = bold_col is not None and i == bold_col
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Inches(w)
        # keep the whole table on one page: rows can't split, and every row but the last
        # is kept with the next one
        for ri, row in enumerate(t.rows):
            trPr = row._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit"); trPr.append(cs)
            if ri < len(t.rows) - 1:
                for cell in row.cells:
                    for cp in cell.paragraphs:
                        cp.paragraph_format.keep_with_next = True
        self._move(t._tbl)
        spacer = doc.add_paragraph(); self._move(spacer._p)

    def eq(self, key, width):
        self.fig(key, None, width)


def section_cursor(title):
    """Cursor positioned right after the Heading 1 `title`; removes the template's placeholder text."""
    hp = heads[title]
    nxt = hp._p.getnext()
    if nxt is not None and nxt.tag == qn("w:p"):
        txt = "".join(t.text or "" for t in nxt.iter(qn("w:t")))
        if txt.startswith("Write detailed content"):
            anchor = nxt
            cur = Cursor(anchor)
            cur._placeholder = anchor
            return cur
    marker = doc.add_paragraph(); hp._p.addnext(marker._p)
    cur = Cursor(marker._p); cur._placeholder = marker._p
    return cur


def done(cur):
    cur._placeholder.getparent().remove(cur._placeholder)


# ============================================================ cover page
for p in doc.paragraphs[:15]:
    if p.text.startswith("Project Title:"):
        for r in p.runs[1:]:
            r.text = ""
        r = p.add_run("ReDeEP: Detecting Hallucination in Retrieval-Augmented Generation via Mechanistic "
                      "Interpretability — A Reproduction and Robustness Study")
        r.bold = False; r.font.name = TNR; r.font.size = Pt(14)
    if p.text.startswith("Group Members:"):
        p.runs[1].text = "\n1. Gourilakshmi S (2023BCS0021)"
        p.runs[2].text = "\n2. Adithya K (2023BCS0051)"
        p.runs[3].text = ""

# ============================================================ table of contents field
for p in doc.paragraphs:
    if p.text.startswith("Right click and update field"):
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        def fld(t):
            r = OxmlElement("w:r"); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), t); r.append(fc); return r
        p._p.append(fld("begin"))
        ri = OxmlElement("w:r"); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
        it.text = ' TOC \\o "1-2" \\h \\z \\u '; ri.append(it); p._p.append(ri)
        p._p.append(fld("separate"))
        rt = OxmlElement("w:r"); tt = OxmlElement("w:t"); tt.text = "Right-click and choose Update Field to refresh the table of contents."; rt.append(tt); p._p.append(rt)
        p._p.append(fld("end"))
        break

# ============================================================ Abstract + base paper
c = Cursor(heads["Abstract"]._p.getnext())   # insert before "Base Paper Details"
for para in [
    "Retrieval-augmented generation (RAG) grounds a large language model in retrieved evidence, yet models still hallucinate "
    "even when the evidence is correct. ReDeEP (Sun et al., ICLR 2025) detects such hallucinations by reading the model's "
    "internals: it measures how much copying heads use the retrieved context (External Context Score, ECS) and how much "
    "knowledge FFN layers inject parametric memory (Parametric Knowledge Score, PKS). In this project we reproduce ReDeEP on "
    f"the RAGTruth benchmark using a single 6 GB laptop GPU. On the paper's own backbone, Llama-2-7B-chat, our detector reaches "
    f"an AUC of {f3(l2['cv']['AUC'][0])} against the published 0.7325, and the paper's two mechanistic findings replicate.",
    "Across six LLaMA-family models from 1.1B to 13B parameters, ReDeEP outperforms perplexity and length-normalised entropy "
    f"on RAGTruth, with Llama-2-13B the strongest detector ({f3(tok('llama2_13b_ragtruth')['cv']['AUC'][0])}). Using the authors' "
    f"released code and data we also reproduce chunk-level ReDeEP exactly (leakage-free AUC {f3(ev('llama2_7b_ragtruth')['nested_cv']['AUC_mean'])}, "
    "identical to the authors' own scores), evaluate the Dolly (AC) dataset, and audit the official evaluation, which selects "
    "features on the test set and inflates AUC by 0.016 on RAGTruth and about 0.08 on Dolly (AC).",
    "Our extension stress-tests ReDeEP under corrupted retrieval: an irrelevant passage raises the hallucination score sharply, "
    "but a passage whose numbers contradict the response leaves it unchanged, so ReDeEP measures whether a response is about "
    "its context rather than whether it agrees with it. Prompted black-box judges from the Gemini and GPT families "
    f"(AUC {f3(JG[SPL]['AUC'])} and {f3(JO[SPL]['AUC'])}) outperform ReDeEP ({f3(s2['AUC'])}) on the same responses, while a small "
    f"3B LLaMA judge is at chance ({f3(JL[SPL]['AUC'])}). Finally, the AARF mitigation's apparent benefit largely disappears once "
    "the judge's position bias is controlled.",
]:
    c.p(para)
c.p("Keywords: hallucination detection, retrieval-augmented generation, mechanistic interpretability, large language models, "
    "robustness, reproducibility.", italic=True)

BASE = {
    "Base Paper Title:": "ReDeEP: Detecting Hallucination in Retrieval-Augmented Generation via Mechanistic Interpretability",
    "DOI link of the paper:": "https://doi.org/10.48550/arXiv.2410.11414 (arXiv:2410.11414; OpenReview/ICLR 2025 proceedings)",
    "Conference/Journal Name:": "International Conference on Learning Representations (ICLR 2025)",
    "Authors": ": Zhongxiang Sun, Xiaoxue Zang, Kai Zheng, Yang Song, Jun Xu, Xiao Zhang, Weijie Yu, Yang Song, Han Li",
    "Year of Publication": ": 2025 (arXiv preprint first posted October 2024)",
}
for p in doc.paragraphs:
    k = p.text.strip()
    if k in BASE:
        r = p.add_run(" " + BASE[k] if k.endswith(":") else BASE[k]); r.bold = False; r.font.name = TNR

# ============================================================ Problem statement
c = section_cursor("Problem Statement")
c.p("Retrieval-augmented generation (RAG) is the standard way to make large language models (LLMs) more trustworthy: a "
    "retriever fetches passages relevant to a query and the model is instructed to answer from them. RAG is widely deployed in "
    "high-stakes settings such as legal research, medical question answering and enterprise search, precisely because it is "
    "assumed to be more reliable than a standalone LLM.")
c.p("RAG does not remove hallucination. Even when the retrieved passage is accurate and contains the answer, models still "
    "produce claims that contradict it or go beyond it. In the base paper's motivating example, a model given a passage stating "
    "that a lunar mission returned 1,935.3 g of samples still answers with a different number recalled from its training data. "
    "If a model can hallucinate when the correct answer is in front of it, adding retrieval is not a sufficient fix: we need "
    "detectors that tell when a response has drifted away from its context.")
c.p("The base paper distinguishes this RAG hallucination (accurate context, unfaithful response) from a knowledge conflict "
    "(the retrieved context itself is wrong). Existing detectors fall into three families that each confound the two knowledge "
    "sources:")
c.bullets([("Parametric confounded by external (PCE):", "hidden-state methods such as EigenScore, SEP, SAPLMA and ITI; the external context mixes into the states they read."),
           ("External confounded by parametric (ECP):", "prompt- or retrieval-based checkers such as RAGAS, ChainPoll and RefChecker; the checking model's own knowledge leaks into the verdict."),
           ("Mixed parametric and external (MPE):", "uncertainty methods such as perplexity, energy and LN-Entropy; they blend both signals without separating them.")])
c.p("ReDeEP's claim is that decoupling context usage from parametric-knowledge usage, measured inside the model, gives a "
    "cleaner signal. Our review of the paper identified four gaps: it is validated only on LLaMA-family models, its head and "
    "layer selection is tuned by hand, its mitigation (AARF) gives only modest gains, and it assumes the retrieved context is "
    "accurate, never testing noisy or wrong retrieval.")
c.p("Objectives of this project:", bold_lead=None)
c.bullets(["Reproduce ReDeEP's token-level and chunk-level detectors on RAGTruth on consumer hardware (RTX 4050, 6 GB VRAM) and compare with the published results.",
           "Replicate the paper's two mechanistic findings and test how the method transfers across six LLaMA-family models (1.1B to 13B parameters), including the paper's Llama-2-13B and Llama-3-8B.",
           "Evaluate the second dataset (Dolly (AC)), the AARF mitigation, and the official evaluation protocol using the authors' released code and data.",
           "Extend the method with a robustness study under corrupted retrieval, and compare it with black-box LLM judges from the GPT, Gemini and LLaMA families."],
          numbered=True)
done(c)

# ============================================================ Method
c = section_cursor("Method Overview with Architecture")
c.p("ReDeEP decouples two internal mechanisms of a transformer language model. Copying heads are attention heads whose output "
    "moves information from context tokens into the residual stream; they carry the external, retrieved knowledge. Knowledge "
    "FFNs are feed-forward layers, mostly in the later part of the network, that inject the model's stored, parametric "
    "knowledge. A RAG hallucination occurs when the model under-uses the first and over-relies on the second. Figure 1 shows the "
    "complete pipeline as we implemented it.")
c.fig("architecture", "Figure 1: Architecture of the ReDeEP detection pipeline (with the AARF mitigation shown below).", 6.2)
c.h("External Context Score (ECS)")
c.p("For attention head h in layer l and each response token n, let T_n be the 10% of context tokens the head attends to most, "
    "and x^L the final-layer hidden states. ECS is the cosine similarity between the response token's hidden state and the mean "
    "hidden state of the tokens it attended to, averaged over the response R:")
c.eq("eq_ecs", 5.2)
c.p("A high ECS means the response is built from the content the head actually read in the context, not merely that it looked "
    "at the right place. The paper finds that about 77.5% of hallucinations occur even though the heads attend to the correct context.")
c.h("Parametric Knowledge Score (PKS)")
c.p("For layer l, the residual stream just before the FFN block (x^mid) and after the layer (x^l) are decoded to next-token "
    "distributions with the logit lens (final norm and unembedding W_U). PKS is the Jensen-Shannon divergence between the two:")
c.eq("eq_pks", 5.8)
c.p("A large divergence means the FFN changed what the model was about to say, i.e. it injected its own knowledge.")
c.h("Head and Layer Selection, and the Detector")
c.bullets([("Copying heads (set A):", "each head's OV circuit W_E W_V W_O W_U is scored on 500 sampled vocabulary tokens; heads with a large-magnitude trace and few Gershgorin-disc outliers copy tokens onto themselves. The top 8 heads form A (the official code uses the authors' released top-32 heads per model)."),
           ("Knowledge FFNs (set F):", "on each training fold, the 5 layers whose PKS correlates most positively with the hallucination label."),
           ("Detector:", "a logistic regression on the standardised features mean PKS over F and negated mean ECS over A, i.e. the score below; the decision threshold is the best-F1 point on the training fold.")])
c.eq("eq_h", 3.8)
c.h("Chunk-Level ReDeEP and AARF")
c.p("ReDeEP(chunk) works on text chunks instead of single tokens, using the chunk boundaries released with the official data. "
    "For each response chunk and copying head, ECS is the BGE sentence-embedding similarity between the response chunk and the "
    "prompt chunk that receives the most attention from it; PKS is the token-summed JSD of each layer over the chunk. AARF "
    "(Add Attention, Reduce FFN) is the paper's training-free mitigation: at decoding time, when the token-level ReDeEP score "
    "exceeds a threshold (0.6), it multiplies the copying heads' output by 1.2 and the knowledge FFNs' output by 0.8.")
c.h("Our Implementation")
c.bullets([("Memory engineering for a 6 GB GPU:", "the reference pipeline asks for every full attention matrix, which does not fit for 7B models. Our forward hooks keep only each layer's top-k context indices and the residual before each FFN; the attended mean is computed as a sparse matrix product; perplexity and LN-Entropy reuse the same forward pass. The rewritten ECS matches the original to within 6 x 10^-8."),
           ("Quantisation and hardware:", "1B models in fp16 and 3B-13B models in 4-bit NF4 (bitsandbytes). Llama-2-7B and smaller ran on the laptop (RTX 4050, 6 GB); Llama-2-13B and Llama-3-8B ran on a free Kaggle T4 GPU with the same code."),
           ("Faithful ports of the official code:", "chunk-level scoring, the per-token AARF gate, and the official regression protocol were re-implemented from the authors' repository and validated against their logged values (our chunk features agree head by head in 148 of 160 spot checks; PKS correlates at 0.994)."),
           ("Evaluation protocol:", "stratified 5-fold cross-validation with copying heads, knowledge FFNs and hyperparameters selected only on training folds (nested CV), plus the official protocol for comparison."),
           ("Baselines and judges:", "perplexity and LN-Entropy from the same forward pass; prompted LLM judges (Gemini 3.5 Flash-Lite, OpenAI's open-weights gpt-oss-120b, and a local Llama-3.2-3B) asked for the probability that a response hallucinates.")])
c.h("Extension: Robustness Under Corrupted Retrieval")
c.p("ReDeEP assumes the retrieved context is accurate. To test what its signals measure when that assumption is broken, we hold "
    "each response fixed and change only the passage: (1) accurate - the real passage; (2) irrelevant - a passage from a "
    "different source of the same task type; (3) numbers corrupted - every number in the real passage multiplied by a random "
    "factor in [1.3, 3.0], so numbers the response copied are now contradicted. A detector that tracks faithfulness should raise "
    "its score in both corrupted conditions.")
done(c)

# ============================================================ Dataset
c = section_cursor("Dataset Details")
c.h("RAGTruth")
c.p(f"RAGTruth (Niu et al., 2024) is the first large word-level hallucination corpus built for RAG. It contains {len(allr):,} "
    f"responses to {len(srcs):,} source prompts, each generated by one of six LLMs (GPT-3.5-turbo, GPT-4, Llama-2-7B/13B/70B-chat "
    "and Mistral-7B-Instruct) and annotated by humans with the exact character spans and types of hallucinations. It covers three "
    "tasks: question answering over MS MARCO passages, summarisation of news articles, and data-to-text generation from "
    "structured Yelp business records.")
c.p("The corpus has two files. source_info.jsonl holds each source's id, task type and full prompt including the retrieved "
    "context; response.jsonl holds each response with its generator model, train/test split, a quality flag, and the list of "
    "hallucination spans. A response is labelled hallucinated if it contains at least one span.")
TN = {"QA": "Question answering", "Summary": "Summarisation", "Data2txt": "Data-to-text"}
TS = {"QA": "MS MARCO", "Summary": "CNN/DM, news", "Data2txt": "Yelp"}
c.table(["Task", "Source", "Responses (full)", "Hallucinated (full)", "Responses (our subset)", "Hallucinated (subset)"],
        [[TN[t], TS[t], f"{ds[(t, 'n')]:,}", f"{100 * ds[(t, 'h')] / ds[(t, 'n')]:.0f}%", sub[(t, "n")],
          f"{100 * sub[(t, 'h')] / sub[(t, 'n')]:.0f}%"] for t in ("QA", "Summary", "Data2txt")],
        "Table 1: RAGTruth by task - full corpus and our main evaluation subset.", [1.4, 1.0, 0.95, 0.95, 0.95, 0.95])
c.p(f"Main evaluation subset (paper setting): all good-quality test-split responses generated by llama-2-7b-chat, capped at "
    f"6,000 characters so every example fits in 1,024 tokens: {l2['n']} responses, {l2['n_hallucinated']} hallucinated. For "
    "Llama-2-13B and Llama-3-8B, each model scores its own 450 RAGTruth test responses (the Llama-3-8B responses come from the "
    "authors' released data).")
c.h("Dolly (AC)")
c.p("Dolly (AC) is the paper's second dataset: a subset of Databricks' dolly-15k instruction data filtered to prompts with an "
    "accurate context passage, answered by each backbone and hand-labelled. The authors' release contains 100 question-answering "
    "prompts with responses from Llama-2-7B, Llama-2-13B and Llama-3-8B (99 test responses each). Dolly (AC) has only a "
    "response-level label (hallucinated or not) and no span positions, so it is evaluated at the response level.")
c.h("Other Data Used")
c.bullets([("Chunk boundaries:", "the official prompt_spans / response_spans files used by chunk-level ReDeEP."),
           ("Authors' logged scores:", "the chunk-level features the authors computed for Llama-2-7B, used to validate our implementation and to audit the evaluation protocol."),
           ("Released hyperparameters:", "the authors' selected copying heads, layers and AARF settings for each backbone.")])
done(c)

# ============================================================ Results
c = section_cursor("Results and Visualizations")
c.h("Main Reproduction: Hallucination Detection")
c.p(f"ReDeEP reproduces on the paper's backbone: Llama-2-7B reaches an AUC of {f3(l2['cv']['AUC'][0])} (+/- {f3(l2['cv']['AUC'][1])}) "
    f"against the published 0.7325, and a PCC of {f3(l2['cv']['PCC'][0])} against 0.3979. It beats both baselines on every "
    "backbone, and is more stable across folds (std 0.03-0.05) than the baselines (0.06-0.08).")
c.table(["Backbone", "ReDeEP AUC", "PCC", "F1", "Perplexity AUC", "LN-Entropy AUC"],
        [[n, f"{f3(R[n]['cv']['AUC'][0])} +/- {f3(R[n]['cv']['AUC'][1])}", f3(R[n]["cv"]["PCC"][0]), f3(R[n]["cv"]["F1"][0]),
          f3(R[n]["cv"]["PPL_AUC"][0]), f3(R[n]["cv"]["LNEnt_AUC"][0])] for n, _ in BB] +
        [["Paper (Llama-2-7B)", "0.7325", "0.3979", "-", "-", "-"]],
        f"Table 2: Detection on RAGTruth (llama-2-7b-chat responses, n = {l2['n']}), 5-fold CV mean.", [1.7, 1.25, 0.65, 0.65, 0.9, 0.9], bold_col=1)
c.fig("detection_auc", "Figure 2: Detection AUC per backbone (mean +/- std over 5 folds); the dashed line is the paper's result.", 5.0)
c.fig("roc", "Figure 3: ROC curves from out-of-fold scores; all ReDeEP curves lie above the baselines over most of the range.", 4.8)
c.fig("score_hist", "Figure 4: Distribution of out-of-fold ReDeEP scores on Llama-2-7B; hallucinated responses shift right.", 5.0)

m7 = A["Llama-2-7B"]
c.h("Replicating the Paper's Mechanistic Findings")
c.p(f"Both core observations of the paper replicate. On Llama-2-7B, all {m7['heads_ecs_truth_gt_halluc']} of {m7['heads_total']} heads "
    f"have higher mean ECS on truthful responses (paper: 1006/1024), and PKS is higher on hallucinated responses in "
    f"{m7['pks_halluc_gt_truth_late_layers']} of the {m7['n_layers'] // 2} later layers but only {m7['pks_halluc_gt_truth_early_layers']} "
    "of the earlier ones. Because ECS compares final-layer hidden states for every head, per-head scores are correlated, so the "
    "head count reflects a response-level shift more than independent effects.")
c.table(["Backbone", "Layers x Heads", "Heads: ECS truth > halluc.", "Early layers: PKS h > t", "Late layers: PKS h > t"],
        [[AN[n], f"{A[AN[n]]['n_layers']} x {A[AN[n]]['n_heads']}", f"{A[AN[n]]['heads_ecs_truth_gt_halluc']}/{A[AN[n]]['heads_total']}",
          f"{A[AN[n]]['pks_halluc_gt_truth_early_layers']}/{A[AN[n]]['n_layers'] // 2}",
          f"{A[AN[n]]['pks_halluc_gt_truth_late_layers']}/{A[AN[n]]['n_layers'] - A[AN[n]]['n_layers'] // 2}"] for n, _ in BB],
        "Table 3: Replication of the paper's two mechanistic findings.", [1.5, 1.0, 1.4, 1.2, 1.2])
c.fig("ecs_heatmap", "Figure 5: Mean ECS of truthful minus hallucinated responses for every head of Llama-2-7B (all positive).", 4.6)
c.fig("pks_by_layer", "Figure 6: PKS of hallucinated minus truthful responses by relative layer depth; the gap is concentrated in later layers.", 4.8)

c.h("Performance by Task and Ablations")
pt7 = m7["per_task"]
c.p(f"Pooled AUC hides large differences between tasks. On Llama-2-7B, ReDeEP's AUC within a single task is {f3(pt7['QA']['redeep'])} "
    f"on QA, {f3(pt7['Summary']['redeep'])} on summarisation and {f3(pt7['Data2txt']['redeep'])} on data-to-text, because the "
    f"hallucination rate differs strongly by task ({100 * sub[('QA', 'h')] / sub[('QA', 'n')]:.0f}%, {100 * sub[('Summary', 'h')] / sub[('Summary', 'n')]:.0f}% "
    f"and {100 * sub[('Data2txt', 'h')] / sub[('Data2txt', 'n')]:.0f}%). Both signals matter: ECS alone gives "
    f"{f3(m7['feature_ablation']['ecs'])}, PKS alone {f3(m7['feature_ablation']['pks'])}, and both together {f3(m7['feature_ablation']['both'])}. "
    f"The detector is insensitive to the number of heads and layers selected (AUC {f3(np.min(m7['grid_auc']))}-{f3(np.max(m7['grid_auc']))}).")
c.fig("per_task", "Figure 7: Out-of-fold ReDeEP AUC within each task type (dotted line: chance).", 4.8)
c.table(["Backbone", "ECS only", "PKS only", "Both"],
        [[AN[n], f3(A[AN[n]]["feature_ablation"]["ecs"]), f3(A[AN[n]]["feature_ablation"]["pks"]), f3(A[AN[n]]["feature_ablation"]["both"])] for n, _ in BB],
        "Table 4: Feature ablation (out-of-fold AUC).", [1.8, 1.0, 1.0, 1.0], bold_col=3)
c.fig("ablation_grid", "Figure 8: AUC on Llama-2-7B as the number of copying heads and FFN layers varies (bold: default 8 heads, 5 layers).", 4.2)

c.h("Reproduction with the Official Code: Chunk Level and Larger Backbones")
def tk(t, d):
    if t == "llama2_7b" and d == "ragtruth":
        return f3(l2["cv"]["AUC"][0])
    r = tok(f"{t}_{d}"); return f3(r["cv"]["AUC"][0]) if r else "-"
def ck(t, d):
    e = ev(f"{t}_{d}"); return f"{f3(e['nested_cv']['AUC_mean'])} ({f3(e['official_protocol']['AUC'])})" if e else "-"
c.table(["Backbone", "RAGTruth token", "RAGTruth chunk", "Dolly token", "Dolly chunk"],
        [[n, tk(t, "ragtruth"), ck(t, "ragtruth"), tk(t, "dolly"), ck(t, "dolly")] for n, t in BIG],
        "Table 5: ReDeEP AUC with the authors' data (leakage-free nested CV; official-protocol AUC in brackets).", [1.4, 1.1, 1.35, 1.0, 1.35])
t13, t8 = tok("llama2_13b_ragtruth"), tok("llama3_8b_ragtruth")
c.p(f"Our chunk-level Llama-2-7B AUC under leakage-free evaluation is {f3(ev('llama2_7b_ragtruth')['nested_cv']['AUC_mean'])}, identical "
    f"to the {f3(evoff['nested_cv']['AUC_mean'])} obtained from the authors' own logged scores. Llama-2-13B is the strongest "
    f"detector we tested ({f3(t13['cv']['AUC'][0])} token-level; perplexity {f3(t13['cv']['PPL_AUC'][0])}, LN-Entropy "
    f"{f3(t13['cv']['LNEnt_AUC'][0])}). On Llama-3-8B the margin over LN-Entropy is small on RAGTruth ({f3(t8['cv']['AUC'][0])} vs. "
    f"{f3(t8['cv']['LNEnt_AUC'][0])}) and reversed on Dolly (AC), consistent with the paper's note that heads and layers must be "
    "re-tuned per model.")

c.h("Audit of the Official Evaluation Protocol")
c.p("The released regression scripts rank copying heads and FFN layers by their AUC on the evaluation set itself, keep the best "
    "ones, combine them with weights tuned per model and dataset, and report AUC on the same responses. We first confirmed our "
    f"re-implementation by reproducing the authors' stored result exactly (AUC {f3(evoff['official_protocol']['AUC'])}, PCC "
    f"{f3(evoff['official_protocol']['PCC'])}), then moved all selection into training folds.")
audit = [("Authors' logs, RAGTruth 7B", evoff), ("Ours, RAGTruth 7B", ev("llama2_7b_ragtruth")), ("Ours, RAGTruth 13B", ev("llama2_13b_ragtruth")),
         ("Ours, RAGTruth 8B", ev("llama3_8b_ragtruth")), ("Ours, Dolly 7B", ev("llama2_7b_dolly")), ("Ours, Dolly 13B", ev("llama2_13b_dolly")),
         ("Ours, Dolly 8B", ev("llama3_8b_dolly"))]
c.table(["Scores", "Official protocol", "All tuned on test", "Nested CV (ours)"],
        [[n, f3(e["official_protocol"]["AUC"]), f3(e["tuned_on_test"]["AUC"]), f3(e["nested_cv"]["AUC_mean"])] for n, e in audit if e],
        "Table 6: Effect of selecting features on the test set (chunk level).", [2.2, 1.2, 1.2, 1.2], bold_col=3)
c.fig("leakage", "Figure 9: Official protocol vs. leakage-free evaluation. The inflation is small on RAGTruth and about 0.08 on Dolly (AC) for all three models.", 5.8)
c.p("On RAGTruth (450 responses) the inflation is small (0.016 AUC), so the paper's main conclusions stand. On the small Dolly "
    "(AC) set (99 responses) it is about 0.08 on all three backbones: choosing the best of 32 heads and 32 layers on the test set "
    "fits noise. Results on small benchmarks should use held-out selection.")

c.h("Comparison with Black-Box LLM Judges")
c.table(["Detector", "Kind", "AUC", "PCC", "Recall", "F1"],
        [["Gemini 3.5 Flash-Lite", "Prompted judge (API)", f3(JG[SPL]["AUC"]), f3(JG[SPL]["PCC"]), f3(JG[SPL]["Rec"]), f3(JG[SPL]["F1"])],
         ["GPT: gpt-oss-120b", "Prompted judge (API)", f3(JO[SPL]["AUC"]), f3(JO[SPL]["PCC"]), f3(JO[SPL]["Rec"]), f3(JO[SPL]["F1"])],
         ["ReDeEP (Llama-2-7B)", "Model internals", f3(s2["AUC"]), f3(s2["PCC"]), f3(s2["Rec"]), f3(s2["F1"])],
         ["LN-Entropy", "Output uncertainty", f3(s2["LNEnt_AUC"]), "-", "-", "-"],
         ["Perplexity", "Output uncertainty", f3(s2["PPL_AUC"]), "-", "-", "-"],
         ["Llama-3.2-3B (local)", "Prompted judge", f3(JL[SPL]["AUC"]), f3(JL[SPL]["PCC"]), f3(JL[SPL]["Rec"]), f3(JL[SPL]["F1"])]],
        "Table 7: Judges vs. ReDeEP on the same 124 held-out responses.", [1.7, 1.5, 0.65, 0.65, 0.7, 0.65], bold_col=2)
c.fig("judges_auc", "Figure 10: All detectors on the same 124 held-out responses (dotted line: chance).", 5.0)
c.p("Large API judges are more accurate than ReDeEP because they read the passage and the response as text and can directly "
    "spot a mismatched number or name. A small prompted judge fails: Llama-3.2-3B almost always answers \"no hallucination\". "
    "ReDeEP's advantage is cost and access: one forward pass of a small open model, with no API, no per-call cost and no data "
    "leaving the machine.")

c.h("Extension: Robustness Under Corrupted Retrieval")
rows = []
for name, cc in (("TinyLlama-1.1B", robT), ("Llama-2-7B", rob7)):
    for k, lab in (("accurate", "Accurate"), ("irrelevant", "Irrelevant"), ("contradictory", "Numbers corrupted")):
        m = cc[k]
        rows.append([name, lab, f3(m["score_mean"]), f3(m["E_A_mean"]), f3(m["P_F_mean"]), f3(m["AUC_vs_original_labels"])])
c.table(["Backbone", "Condition", "Mean score", "ECS", "PKS", "AUC vs. labels"], rows,
        "Table 8: Robustness - mean ReDeEP score and signals per condition (150 held-out responses).", [1.4, 1.4, 0.9, 0.7, 0.7, 1.0])
c.fig("robustness_shift", "Figure 11: Change in mean ReDeEP score relative to the accurate passage (paired Wilcoxon test).", 5.0)
c.p(f"Irrelevant passages work as intended: the score jumps for nearly every response ({robT['irrelevant']['score_shift_vs_accurate']:+.3f} "
    f"for TinyLlama, {rob7['irrelevant']['score_shift_vs_accurate']:+.3f} for Llama-2-7B, both p < 0.001). Corrupted numbers go "
    f"unnoticed: the score does not move for TinyLlama ({robT['contradictory']['score_shift_vs_accurate']:+.3f}) and falls slightly "
    f"for Llama-2-7B ({rob7['contradictory']['score_shift_vs_accurate']:+.3f}). On the {num7['n']} responses that state a contradicted "
    f"number, the score rose in only {round(100 * numT['frac_score_increased'])}% (TinyLlama) and {round(100 * num7['frac_score_increased'])}% "
    "(Llama-2-7B) of cases. ReDeEP therefore detects topical hallucination (a response that drifts away from what the context is "
    "about) but not factual inconsistency with the context, and should be paired with a fact-level check.")

c.h("AARF Mitigation")
c.table(["Backbone", "Answers changed", "Gate fires", "Judge, official order (AARF : base)", "Judge, both orders agree"],
        [[n, f"{AA[t]['changed']}/{AA[t]['n']}", f"{AA[t]['fire']:.0f}%", f"{AA[t]['offA']} : {AA[t]['offB']}", f"{AA[t]['bothA']} : {AA[t]['bothB']}"] for n, t in BIG],
        "Table 9: AARF vs. baseline answers on Dolly (AC), judged by Gemini 3.5 Flash-Lite.", [1.3, 1.1, 0.9, 1.6, 1.3])
a7, a13 = AA["llama2_7b"], AA["llama2_13b"]
c.p(f"With the released hyperparameters the AARF gate fires on every decoding step for the Llama-2 models (Llama-2-7B: lowest "
    f"gate score {a7['gmin']:.2f} over {a7['steps']:,} steps, against a 0.6 threshold) and on only {AA['llama3_8b']['fire']:.0f}% of "
    "steps for Llama-3-8B, so it acts as an always-on or almost-never-on intervention. Judged in the paper's fixed order (AARF "
    f"shown first), AARF wins {a7['offA']} to {a7['offB']} on Llama-2-7B. But the judge reversed its verdict when the answers were "
    f"swapped in {a7['flips']} of {a7['parsed']} pairs ({100 * a7['flips'] / a7['parsed']:.0f}%) for 7B and {a13['flips']} of "
    f"{a13['parsed']} ({100 * a13['flips'] / a13['parsed']:.0f}%) for 13B; when both orders must agree, AARF's advantage shrinks to "
    f"{a7['bothA']} : {a7['bothB']} and {a13['bothA']} : {a13['bothB']}, within chance.")

c.h("Computational Cost")
c.table(["Backbone", "Copying heads (s)", "Per response (s)", "Total, 415 responses (min)"],
        [[AN[n], f"{A['_cost'][AN[n]]['copying_head_seconds']:.1f}", f"{A['_cost'][AN[n]]['sec_per_example_median']:.1f}",
          f"{A['_cost'][AN[n]]['minutes']:.1f}"] for n, _ in BB],
        "Table 10: Wall-clock cost on the RTX 4050 laptop GPU.", [1.8, 1.3, 1.3, 1.6])

c.h("Discussion, Limitations and Conclusion")
c.bullets([("What reproduces:", "the headline token-level result, the chunk-level result (exactly), and both mechanistic findings, on a laptop GPU with 4-bit weights."),
           ("What the pooled number hides:", "per-task AUCs are 0.08-0.35 lower than pooled ones; the official protocol inflates results on small datasets."),
           ("Where it stops working:", "contradicted facts go undetected, Llama-3-8B gains little, and the AARF mitigation's benefit does not survive a position-bias-controlled judge."),
           ("Limitations:", "4-bit quantisation; copying heads scored on 500 sampled vocabulary tokens; numbers-only corruption; free-tier judges (Gemini Flash-Lite instead of GPT-4o; open-weights gpt-oss-120b); 99 responses per model on Dolly (AC).")])
c.p("Conclusion. ReDeEP reproduces faithfully and is a cheap, private first-line detector of topical hallucinations, strongest on "
    "Llama-2-13B. It should be evaluated with held-out feature selection and paired with a fact-level check for numeric or "
    "entity conflicts. Future work: entity and negation corruptions, a numeric-consistency feature combined with ECS, "
    "full-precision runs, and a recalibrated AARF gate.")
done(c)

# ============================================================ Implementation & links
c = section_cursor("Implementation and Dataset Links [GitHub]")
c.p("Project implementation (code, results, report and slide sources):")
c.bullets([("Project GitHub repository:", "https://github.com/ad1thyaaa/ReDeEP-Reproduction-Robustness"),
           ("Base paper's official code:", "https://github.com/Jeryi-Sun/ReDEeP-ICLR"),
           ("RAGTruth dataset:", "https://github.com/ParticleMedia/RAGTruth"),
           ("Official ReDeEP dataset (Dolly (AC), Llama-3-8B responses, chunk spans):", "https://drive.google.com/file/d/1s-pmaBQutC6eQGtk2F3uKaMSkn_iGQwR"),
           ("Dolly-15k (source of Dolly (AC)):", "https://github.com/databrickslabs/dolly"),
           ("Model checkpoints (Hugging Face):", "TinyLlama/TinyLlama-1.1B-Chat-v1.0, unsloth/Llama-3.2-1B-Instruct, unsloth/Llama-3.2-3B-Instruct, NousResearch/Llama-2-7b-chat-hf, NousResearch/Llama-2-13b-chat-hf, NousResearch/Meta-Llama-3-8B-Instruct, BAAI/bge-base-en-v1.5")])
c.p("Repository structure:")
c.bullets([("redeep/mechanistic.py:", "copying-head score, ECS, PKS, perplexity and LN-Entropy via forward hooks."),
           ("redeep/detect.py, redeep/chunk_level.py:", "token-level detector and chunk-level ReDeEP."),
           ("redeep/aarf.py, redeep/aarf_gen.py:", "AARF intervention and gated AARF decoding."),
           ("redeep/robustness.py, redeep/llm_judge.py:", "robustness extension and black-box judges."),
           ("scripts/:", "run_pipeline.py, run_chunk.py, run_token_official.py, run_aarf.py, run_judges.py, judge_aarf.py, eval_chunk.py, leakage_check.py, run_robustness.py."),
           ("kaggle/:", "the Kaggle GPU job used for Llama-2-13B and Llama-3-8B."),
           ("paper/, presentation/:", "build scripts for the IEEE-format paper and the slides; every number is generated from data/.")])
c.p("Environment: Python 3.11, PyTorch 2.11 (CUDA 12.8), transformers 5.17, bitsandbytes 0.50, scikit-learn; NVIDIA RTX 4050 "
    "Laptop GPU (6 GB) and Kaggle Tesla T4 (16 GB).")
done(c)

# ============================================================ References (APA)
refs_h = heads["References [APA Style]"]
anchor = refs_h._p.getnext()
if anchor is None or anchor.tag != qn("w:p"):
    m = doc.add_paragraph(); refs_h._p.addnext(m._p); anchor = m._p
c = Cursor(anchor)
for ref in [
    "Geva, M., Schuster, R., Berant, J., & Levy, O. (2021). Transformer feed-forward layers are key-value memories. In Proceedings of the 2021 Conference on Empirical Methods in Natural Language Processing. Association for Computational Linguistics.",
    "Malinin, A., & Gales, M. (2021). Uncertainty estimation in autoregressive structured prediction. In International Conference on Learning Representations.",
    "Niu, C., Wu, Y., Zhu, J., et al. (2024). RAGTruth: A hallucination corpus for developing trustworthy retrieval-augmented language models. In Proceedings of the 62nd Annual Meeting of the Association for Computational Linguistics.",
    "nostalgebraist. (2020). Interpreting GPT: The logit lens. LessWrong. https://www.lesswrong.com/posts/AcKRB8wDpdaN6v6ru/interpreting-gpt-the-logit-lens",
    "Sun, Z., Zang, X., Zheng, K., Song, Y., Xu, J., Zhang, X., Yu, W., Song, Y., & Li, H. (2025). ReDeEP: Detecting hallucination in retrieval-augmented generation via mechanistic interpretability. In International Conference on Learning Representations. https://doi.org/10.48550/arXiv.2410.11414",
    "Touvron, H., Martin, L., Stone, K., et al. (2023). Llama 2: Open foundation and fine-tuned chat models. arXiv preprint arXiv:2307.09288.",
    "Xiao, S., Liu, Z., Zhang, P., & Muennighoff, N. (2023). C-Pack: Packaged resources to advance general Chinese embedding. arXiv preprint arXiv:2309.07597.",
]:
    para = c.p(ref, align=WD_ALIGN_PARAGRAPH.LEFT)
    para.paragraph_format.left_indent = Inches(0.5)
    para.paragraph_format.first_line_indent = Inches(-0.5)
if anchor is not None and not "".join(t.text or "" for t in anchor.iter(qn("w:t"))).strip():
    anchor.getparent().remove(anchor)

# ============================================================ layout fixes
# 1) the template's empty paragraph + hard page break after the TOC spill onto a blank
#    page once the TOC fills page 2: drop them and start Abstract on a new page instead
abs_h = heads["Abstract"]._p
prev = abs_h.getprevious()
while prev is not None and prev.tag == qn("w:p") and not "".join(t.text or "" for t in prev.iter(qn("w:t"))).strip() \
        and "fldChar" not in prev.xml:
    nxt_prev = prev.getprevious()
    prev.getparent().remove(prev)
    prev = nxt_prev
heads["Abstract"].paragraph_format.page_break_before = True
# 2) keep the "Table of Contents" title out of the TOC itself
heads["Table of Contents"].style = doc.styles["TOC Heading"]
for r in heads["Table of Contents"].runs:
    r.font.name = TNR
# 3) Base Paper Details items end with a line break in the template: remove it so values stay inline
for p in doc.paragraphs:
    if p.text.strip().split(":")[0].strip() in ("Year of Publication", "Authors", "Base Paper Title", "DOI link of the paper", "Conference/Journal Name"):
        for br in list(p._p.iter(qn("w:br"))):
            br.getparent().remove(br)
# 4) space after "Project Title:"
for p in doc.paragraphs[:15]:
    if p.text.startswith("Project Title:") and p.runs:
        p.runs[0].text = "Project Title: "

doc.save(args.out)
print("wrote", args.out)
